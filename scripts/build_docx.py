#!/usr/bin/env python3
"""
ATS-safe editable DOCX of the Mainframe COBOL Developer resume.

No tables, no text boxes, no images, no headers/footers, no columns.
Right-aligned dates use a right tab stop (native Word feature, parses cleanly).
"""
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import resume_content as C  # noqa: E402

NAVY = RGBColor(0x0F, 0x2A, 0x47)
INK = RGBColor(0x11, 0x14, 0x18)
SLATE = RGBColor(0x3C, 0x4A, 0x59)

S = 1.0                   # global scale, set by auto-fit
BODY = 9.5
FONT = "Calibri"          # universally available -> no substitution surprises in ATS


def set_spacing(p, before=0, after=0, line=1.0):
    pf = p.paragraph_format
    pf.space_before = Pt(before * S)
    pf.space_after = Pt(after * S)
    pf.line_spacing = line


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "8")
    b.set(qn("w:space"), "2")
    b.set(qn("w:color"), "0F2A47")
    bd.append(b)
    pPr.append(bd)


def run(p, text, size=None, bold=False, italic=False, color=INK, font=FONT,
        spacing=None, caps=False):
    if size is None:
        size = BODY
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size * S)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    if caps:
        r.font.all_caps = True
    if spacing:  # character tracking in twentieths of a point
        rPr = r._element.get_or_add_rPr()
        el = OxmlElement("w:spacing")
        el.set(qn("w:val"), str(int(spacing * S * 20)))
        rPr.append(el)
    # ensure east-asian/complex-script mapping too
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), font)
    return r


def sc(v):
    return v * S


def heading(doc, text, first=False):
    p = doc.add_paragraph()
    set_spacing(p, before=0 if first else 6.6, after=3.0, line=1.0)
    run(p, text.upper(), size=10.0, bold=True, color=NAVY, spacing=0.8)
    bottom_border(p)
    return p


def bullet(doc, indent=0.42):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.left_indent = Cm(indent)
    pf.first_line_indent = Cm(-0.30)
    set_spacing(p, before=0, after=1.6, line=1.0)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def tabbed(doc, left, right, size=10.0, lbold=True, width_cm=18.46):
    p = doc.add_paragraph()
    set_spacing(p, before=0, after=0, line=1.0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(width_cm), WD_TAB_ALIGNMENT.RIGHT)
    run(p, left, size=size, bold=lbold, color=INK)
    if right:
        run(p, "\t" + right, size=size * 0.96, bold=True, color=SLATE)
    return p


def build(path, scale=1.0):
    global S
    S = scale
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BODY * S)
    st.font.color.rgb = INK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0

    sec = doc.sections[0]
    sec.top_margin = Cm(1.02)
    sec.bottom_margin = Cm(0.90)
    sec.left_margin = Cm(1.27)      # 0.5in - standard ATS-safe minimum
    sec.right_margin = Cm(1.27)

    # ---------------- header ----------------
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, 0, 2, 1.0)
    run(p, C.NAME, size=21, bold=True, color=NAVY, spacing=1.9)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, 0, 2, 1.0)
    run(p, C.HEADLINE, size=9.2, bold=True, color=SLATE)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, 0, 3, 1.0)
    run(p, "  |  ".join(C.CONTACT), size=8.9, color=INK)
    bottom_border(p)

    # ---------------- summary ----------------
    heading(doc, "Professional Summary", first=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_spacing(p, 0, 0, 1.0)
    run(p, C.SUMMARY)

    # ---------------- experience ----------------
    heading(doc, "Professional Experience")
    for job in C.EXPERIENCE:
        tabbed(doc, job["role"], job["dates"], size=10.3)
        p = doc.add_paragraph()
        set_spacing(p, 1, 3, 1.0)
        run(p, f'{job["company"]}  |  {job["location"]}', size=9.4,
            italic=True, color=SLATE)
        for b in job["bullets"]:
            run(bullet(doc), b)

    # ---------------- projects ----------------
    heading(doc, "Key Projects")
    for pr in C.PROJECTS:
        p = bullet(doc)
        set_spacing(p, 0, 2.0, 1.0)
        run(p, pr["name"], bold=True)
        run(p, " \u2014 " + pr["desc"] + " ")
        run(p, "Technologies:", bold=True)
        run(p, " " + pr["tech"] + ".", italic=True)

    # ---------------- skills ----------------
    heading(doc, "Technical Skills")
    for label, items in C.SKILLS:
        p = bullet(doc)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_spacing(p, 0, 1.5, 1.0)
        run(p, label + ":", bold=True)
        run(p, " " + items)

    # ---------------- achievements ----------------
    heading(doc, "Achievements")
    for title, desc in C.ACHIEVEMENTS:
        p = bullet(doc)
        set_spacing(p, 0, 2.0, 1.0)
        run(p, title, bold=True)
        run(p, " \u2014 " + desc)

    # ---------------- education ----------------
    heading(doc, "Education")
    for i, e in enumerate(C.EDUCATION):
        pp = tabbed(doc, e["degree"], e["dates"], size=9.9)
        if i:
            set_spacing(pp, 3, 0, 1.0)
        p = doc.add_paragraph()
        set_spacing(p, 1, 0, 1.0)
        run(p, e["school"], size=9.4, italic=True, color=SLATE)

    cp = doc.core_properties
    cp.title = "Pratik Shinde - Mainframe COBOL Developer Resume"
    cp.author = "Pratik Shinde"
    cp.subject = ("Mainframe COBOL Developer | HP NonStop (Tandem) | "
                  "Banking Payment Systems | L3 Production Support")
    kw = ("COBOL, Mainframe, HP NonStop, HP Tandem, Guardian, TACL, TAL, ENFORM, DDL, SQL, SQLCI, "
          "DBUX, FUP, Enscribe, SQL/MP, TMF, EMS, OSS, Production Support, Banking, Payments, "
          "SWIFT, SEPA, IBM MQ, RabbitMQ, SFTP, Incident Management, RCA, Java, Spring Boot")
    assert len(kw) <= 255, len(kw)
    cp.keywords = kw

    doc.save(path)
    print("wrote", path)


if __name__ == "__main__":
    out = (sys.argv[1] if len(sys.argv) > 1
           else "Pratik_Shinde_Mainframe_COBOL_Developer_Resume.docx")
    from measure_docx import measure
    lo, hi, best = 0.60, 1.05, None
    for _ in range(28):
        mid = (lo + hi) / 2
        build(out, mid)
        used, avail, _ = measure(out)
        if used <= avail * 0.972:
            best, lo = mid, mid
        else:
            hi = mid
        if hi - lo < 0.001:
            break
    build(out, best or 0.60)
    used, avail, _ = measure(out)
    print(f"docx scale={best:.4f} fill={100*used/avail:.1f}%")
